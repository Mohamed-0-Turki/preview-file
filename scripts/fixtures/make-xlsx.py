"""Generate real .xlsx fixtures with openpyxl for Office View Engine verification."""

import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "/tmp/opencode/fixtures"

THIN = Side(style="thin", color="FF9E9E9E")
MEDIUM = Side(style="medium", color="FF1F4E79")
DOUBLE = Side(style="double", color="FFC0392B")
HAIR = Side(style="hair", color="FFBBBBBB")


def fixture_fidelity() -> str:
    """Styles, merges, widths, heights, formats, borders, fills, alignment, freeze."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_properties.tabColor = "1F4E79"

    headers = ["Region", "Quarter", "Revenue", "Margin %", "Growth", "Notes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF", size=11, name="Calibri")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=MEDIUM, bottom=MEDIUM)

    rows = [
        ("North America", "Q1", 1284500, 0.412, 0.052, "Steady growth across retail and partner channels."),
        ("North America", "Q2", 1398720, 0.428, 0.089, "New partner programme launched in June."),
        ("Europe", "Q1", 987300, 0.376, 0.031, "FX headwind reduced reported revenue by 2%."),
        ("Europe", "Q2", 1042210, 0.394, 0.056, "Offset by pricing changes in DACH."),
        ("Asia Pacific", "Q1", 764900, 0.301, 0.118, "Fastest growing region; new hires in Singapore."),
        ("Asia Pacific", "Q2", 881455, 0.337, 0.152, "APAC margins improved on mix shift."),
        ("Latin America", "Q1", 312480, 0.244, -0.031, "Currency volatility reduced reported growth."),
        ("Latin America", "Q2", 298110, 0.219, -0.046, "Restructuring completed in July."),
    ]
    for row in rows:
        ws.append(row)

    money = '#,##0.00'
    pct = '0.0%'
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=3).number_format = money
        ws.cell(row=r, column=4).number_format = pct
        ws.cell(row=r, column=5).number_format = '+0.0%;-0.0%;0.0%'
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
        ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")

    ws.merge_cells("A1:F1")
    ws["A1"] = "Quarterly revenue report"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")

    ws.merge_cells("A11:C11")
    ws["A11"] = "Merged banner spanning three columns"
    ws["A11"].font = Font(italic=True, size=12)
    ws["A11"].alignment = Alignment(horizontal="center")
    ws["A11"].fill = PatternFill("solid", fgColor="DDEBF7")
    ws["A11"].border = Border(left=THIN, right=THIN, top=THIN, bottom=DOUBLE)

    ws.merge_cells("A12:B13")
    ws["A12"] = "Two-by-two merge"
    ws["A12"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A12"].fill = PatternFill("solid", fgColor="E2EFDA")

    widths = {"A": 18, "B": 10, "C": 15, "D": 11, "E": 11, "F": 46}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[11].height = 22
    ws.row_dimensions[12].height = 18
    for r in range(2, 10):
        ws.row_dimensions[r].height = 30

    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 100

    ws.conditional_formatting.add(
        "E2:E9",
        CellIsRule(operator="greaterThan", formula=["0.08"], font=Font(bold=True, color="FF1E7B34"),
                   fill=PatternFill("solid", fgColor="FFE2EFDA")),
    )
    ws.conditional_formatting.add(
        "E2:E9",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="FFC0392B")),
    )
    ws.conditional_formatting.add(
        "D2:D9",
        ColorScaleRule(start_type="min", start_color="FFF8696B",
                       mid_type="percentile", mid_value=50, mid_color="FFFFEB84",
                       end_type="max", end_color="FF63BE7B"),
    )
    ws.conditional_formatting.add(
        "C2:C9",
        DataBarRule(start_type="min", end_type="max", color="FF638EC6", showValue=True),
    )
    ws.conditional_formatting.add(
        "G2:G9",
        IconSetRule("3TrafficLights1", "percent", [0, 33, 67], showValue=True),
    )
    ws["G1"] = "Score"
    for i, r in enumerate(range(2, 10), start=1):
        ws.cell(row=r, column=7).value = i * 13

    dv = DataValidation(type="list", formula1='"Draft,Review,Final"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H9")
    ws["H1"] = "Status"

    # Pattern fill + rotation + indent showcase.
    ws["I1"] = "Patterns"
    patterns = [
        ("solid", "FFD9E1F2"),
        ("darkGray", "FF7F7F7F"),
        ("mediumGray", "FFBFBFBF"),
        ("lightGray", "FFD9D9D9"),
        ("darkHorizontal", "FF808080"),
        ("lightVertical", "FFBFBFBF"),
        ("lightDown", "FFA6A6A6"),
        ("gray125", "FFEDEDED"),
        ("gray0625", "FFF7F7F7"),
    ]
    for i, (pattern, color) in enumerate(patterns, start=2):
        cell = ws.cell(row=i, column=9)
        cell.value = pattern
        cell.fill = PatternFill(patternType=pattern, fgColor=color, bgColor="FFFFFFFF")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws["K1"] = "Rotated / indented"
    ws["K1"].font = Font(bold=True)
    for i, text in enumerate(("rotated 45", "rotated 90", "rotated 135"), start=2):
        cell = ws.cell(row=i, column=11)
        cell.value = text
        cell.alignment = Alignment(textRotation=45 * (i - 1), indent=2, vertical="bottom")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    return _save(wb, "fidelity.xlsx")


def fixture_charts() -> str:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["Month", "Series A", "Series B", "Series C"])
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug"]
    values = {
        "Series A": [42, 51, 39, 66, 72, 61, 84, 79],
        "Series B": [31, 45, 52, 48, 63, 70, 58, 74],
        "Series C": [18, 22, 35, 29, 41, 38, 52, 61],
    }
    for i, month in enumerate(months, start=2):
        data.cell(row=i, column=1).value = month
        for j, key in enumerate(("Series A", "Series B", "Series C"), start=2):
            data.cell(row=i, column=j).value = values[key][i - 2]
    for col in "ABCD":
        data.column_dimensions[col].width = 14

    charts = wb.create_sheet("Charts")
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Monthly revenue by series"
    bar.y_axis.title = "Units"
    bar.x_axis.title = "Month"
    bar.height = 8
    bar.width = 16
    bar.add_data(Reference(data, min_col=2, max_col=4, min_row=1, max_row=9), titles_from_data=True)
    bar.set_categories(Reference(data, min_col=1, min_row=2, max_row=9))
    charts.add_chart(bar, "B2")

    stacked = BarChart()
    stacked.type = "bar"
    stacked.grouping = "stacked"
    stacked.overlap = 100
    stacked.title = "Stacked horizontal"
    stacked.height = 8
    stacked.width = 16
    stacked.add_data(Reference(data, min_col=2, max_col=4, min_row=1, max_row=9), titles_from_data=True)
    stacked.set_categories(Reference(data, min_col=1, min_row=2, max_row=9))
    charts.add_chart(stacked, "B20")

    line = LineChart()
    line.title = "Trend line"
    line.height = 8
    line.width = 16
    line.add_data(Reference(data, min_col=2, max_col=4, min_row=1, max_row=9), titles_from_data=True)
    line.set_categories(Reference(data, min_col=1, min_row=2, max_row=9))
    charts.add_chart(line, "B38")

    pie = PieChart()
    pie.title = "Series A share"
    pie.height = 8
    pie.width = 16
    pie.add_data(Reference(data, min_col=2, min_row=1, max_row=9), titles_from_data=True)
    pie.set_categories(Reference(data, min_col=1, min_row=2, max_row=9))
    charts.add_chart(pie, "B56")

    return _save(wb, "charts.xlsx")


def fixture_grid() -> str:
    """Wide + tall sheet: virtualisation, hidden rows/cols, outline, number formats."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    ws.append(["id"] + [f"c{i}" for i in range(1, 41)])
    for r in range(2, 602):
        ws.cell(row=r, column=1).value = r - 1
        for c in range(2, 42):
            value = (r - 1) * c
            cell = ws.cell(row=r, column=c)
            if c % 7 == 0:
                cell.value = f"text-{r}-{c}"
            else:
                cell.value = value
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "B2"
    for col in range(2, 42):
        ws.column_dimensions[get_column_letter(col)].width = 11
    for r in (50, 100, 150, 200, 250, 300, 350, 400, 450, 500):
        ws.row_dimensions[r].outlineLevel = 1
    ws.row_dimensions[600].hidden = True
    ws.column_dimensions["AI"].hidden = True
    return _save(wb, "large-grid.xlsx")


def fixture_types() -> str:
    """Date/currency/percent/scientific/text/error/boolean number formats."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Types"
    import datetime

    ws.append(["kind", "value", "formatted"])
    samples = [
        ("date", datetime.datetime(2026, 3, 14, 15, 9, 26), "yyyy-mm-dd hh:mm:ss"),
        ("date-short", datetime.date(2026, 12, 25), "d mmm yyyy"),
        ("time", datetime.time(13, 45, 0), "h:mm:ss AM/PM"),
        ("currency", 1234567.891, "$#,##0.00"),
        ("percent", 0.4237, "0.00%"),
        ("scientific", 12345.6789, "0.000E+00"),
        ("fraction", 0.666666, "# ?/?"),
        ("text", "00042", "@"),
        ("accounting", -98765.43, "_-$* #,##0.00_-;-$* (#,##0.00)"),
        ("boolean", True, "General"),
        ("error", "#DIV/0!", "General"),
    ]
    for i, (kind, value, fmt) in enumerate(samples, start=2):
        ws.cell(row=i, column=1).value = kind
        ws.cell(row=i, column=2).value = value
        ws.cell(row=i, column=2).number_format = fmt
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 4
    return _save(wb, "types.xlsx")


def _save(wb, name) -> str:
    import os

    os.makedirs(OUT, exist_ok=True)
    path = f"{OUT}/{name}"
    wb.save(path)
    print(f"wrote {path} ({os.path.getsize(path)} bytes)")
    return path


if __name__ == "__main__":
    fixture_fidelity()
    fixture_charts()
    fixture_grid()
    fixture_types()
